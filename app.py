import os
import re
import hmac
from io import BytesIO
from datetime import datetime

import pandas as pd
from flask import Flask, abort, jsonify, render_template, request, send_file, flash, redirect, url_for, session
from werkzeug.utils import secure_filename
from werkzeug.middleware.dispatcher import DispatcherMiddleware
from openpyxl.styles import Font, Border, Side
from openpyxl.utils import get_column_letter
from blob_storage import delete_blob, download_private_blob

app = Flask(__name__, static_folder="public/static")
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "local-development-only-secret")
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(os.environ.get("VERCEL")),
)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB


def blob_uploads_enabled():
    return bool(os.environ.get("VERCEL") and os.environ.get("BLOB_READ_WRITE_TOKEN"))


@app.before_request
def require_shared_password():
    if os.environ.get("VERCEL"):
        if not os.environ.get("APP_ACCESS_PASSWORD"):
            abort(503, description="Set APP_ACCESS_PASSWORD in the Vercel project environment.")
        if not os.environ.get("FLASK_SECRET_KEY"):
            abort(503, description="Set FLASK_SECRET_KEY in the Vercel project environment.")
        if not blob_uploads_enabled():
            abort(503, description="Connect a Vercel Blob store and set BLOB_READ_WRITE_TOKEN.")
    if not os.environ.get("APP_ACCESS_PASSWORD"):
        return
    if request.endpoint in {"login", "upload_auth"} or request.path.startswith("/static/"):
        return
    if session.get("authenticated"):
        return
    if request.path == "/process" or request.is_json:
        return jsonify({"error": "Sign in before using the attendance tools."}), 401
    return redirect(url_for("login"))


@app.context_processor
def inject_blob_upload_setting():
    return {"blob_uploads_enabled": blob_uploads_enabled()}


@app.route("/login", methods=["GET", "POST"])
def login():
    configured_password = os.environ.get("APP_ACCESS_PASSWORD")
    if not configured_password:
        abort(503, description="Set APP_ACCESS_PASSWORD to enable sign-in.")

    error = None
    if request.method == "POST":
        supplied_password = request.form.get("password", "")
        if hmac.compare_digest(supplied_password, configured_password):
            session["authenticated"] = True
            return redirect(url_for("index"))
        error = "The password was not accepted."

    return render_template("login.html", error=error)


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/__upload-auth")
def upload_auth():
    if session.get("authenticated") and os.environ.get("APP_ACCESS_PASSWORD"):
        return "", 204
    return "", 401

REQUIRED_COLUMNS = [
    "HTNO", "SUBJECTCODE", "SUBJECTNAME", "SUBJECTTYPE",
    "FORTNIGHT", "FORTNIGHT_START_DATE", "FORTNIGHT_END_DATE"
]
COLUMN_ALIASES = {
    "FROM_DATE": "FORTNIGHT_START_DATE",
    "TO_DATE": "FORTNIGHT_END_DATE",
}
OUTPUT_COLUMNS = [
    "DEGREE", "COLLEGECODE", "HTNO", "BRANCHCODE", "SECTION", "REGULATION",
    "YEAR", "SEMESTER", "SUBJECTCODE", "SUBJECTNAME", "SUBJECTTYPE",
    "FORTNIGHT", "CLASSESCONDUCTED", "CLASSESATTENDED",
    "FORTNIGHT_START_DATE", "FORTNIGHT_END_DATE"
]
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}  # .xls is intentionally not supported by openpyxl


def allowed_file(filename):
    return os.path.splitext(filename.lower())[1] in ALLOWED_EXTENSIONS


def safe_sheet_name(value, used):
    """Excel worksheet names: <=31 chars, no : \\ / ? * [ ] characters, unique."""
    cleaned = re.sub(r'[:\\/?*\[\]]', "-", str(value)).strip("'").strip()
    cleaned = cleaned or "Subject"
    cleaned = cleaned[:31]
    candidate = cleaned
    counter = 2
    while candidate.lower() in used:
        suffix = f"-{counter}"
        candidate = cleaned[:31-len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


def excel_date(value):
    if pd.isna(value) or not str(value).strip():
        return None
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    return parsed.to_pydatetime() if pd.notna(parsed) else str(value).strip()


def make_output(dataframe, details):
    output = BytesIO()
    used_sheet_names = set()
    student_numbers = dataframe["HTNO"].drop_duplicates().tolist()
    period_columns = ["FORTNIGHT", "FORTNIGHT_START_DATE", "FORTNIGHT_END_DATE"]
    periods = [
        tuple(str(value).strip() if pd.notna(value) else "" for value in row)
        for row in dataframe[period_columns].drop_duplicates().itertuples(index=False, name=None)
        if str(row[0]).strip()
    ]

    def period_sort_key(period):
        try:
            fortnight = (0, float(period[0]))
        except ValueError:
            fortnight = (1, period[0].casefold())
        return fortnight + (period[1], period[2])

    periods.sort(key=period_sort_key)
    output_details = {
        "DEGREE": details["degree"],
        "COLLEGECODE": details["college_code"],
        "BRANCHCODE": details["branch_code"],
        "SECTION": details["section"],
        "REGULATION": details["regulation"],
        "YEAR": details["year"],
        "SEMESTER": details["semester"],
    }
    global_details_by_student = {}
    for row in dataframe.to_dict("records"):
        htno = str(row["HTNO"]).strip()
        student_details = global_details_by_student.setdefault(htno, {})
        for column in output_details:
            value = str(row.get(column, "")).strip() if pd.notna(row.get(column, "")) else ""
            if value and not student_details.get(column):
                student_details[column] = value

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        group_cols = ["SUBJECTCODE", "SUBJECTNAME", "SUBJECTTYPE"]
        grouped = dataframe.groupby(group_cols, dropna=False, sort=False)

        for group_values, group in grouped:
            if any(pd.isna(value) or not str(value).strip() for value in group_values[:2]):
                continue

            subject_rows = group.to_dict("records")
            rows_by_student_period = {}
            subject_details_by_student = {}
            for row in subject_rows:
                htno = str(row["HTNO"]).strip()
                period = tuple(
                    str(row.get(column, "")).strip() if pd.notna(row.get(column, "")) else ""
                    for column in period_columns
                )
                if period[0]:
                    rows_by_student_period.setdefault((htno, period), row)

                student_details = subject_details_by_student.setdefault(htno, {})
                for column in output_details:
                    value = str(row.get(column, "")).strip() if pd.notna(row.get(column, "")) else ""
                    if value and not student_details.get(column):
                        student_details[column] = value

            generated_rows = []
            for htno in student_numbers:
                for period in periods:
                    source_row = rows_by_student_period.get((htno, period), {})
                    student_details = {
                        **global_details_by_student.get(htno, {}),
                        **subject_details_by_student.get(htno, {}),
                    }
                    generated_row = {
                        "HTNO": htno,
                        "SUBJECTCODE": group_values[0],
                        "SUBJECTNAME": group_values[1],
                        "SUBJECTTYPE": group_values[2],
                        "FORTNIGHT": period[0],
                        "FORTNIGHT_START_DATE": period[1],
                        "FORTNIGHT_END_DATE": period[2],
                        "CLASSESCONDUCTED": source_row.get("CLASSESCONDUCTED", ""),
                        "CLASSESATTENDED": source_row.get("CLASSESATTENDED", ""),
                    }
                    for column, fallback in output_details.items():
                        value = source_row.get(column, "")
                        if pd.isna(value) or not str(value).strip():
                            value = student_details.get(column, "")
                        generated_row[column] = str(value).strip() if str(value).strip() else fallback
                    generated_rows.append(generated_row)

            students = pd.DataFrame(generated_rows, columns=OUTPUT_COLUMNS)
            for column in ("FORTNIGHT_START_DATE", "FORTNIGHT_END_DATE"):
                students[column] = students[column].map(excel_date)

            subject_code = str(group_values[0]).strip()
            source_sheets = group["_SOURCE_SHEET"].dropna().unique()
            if len(source_sheets) == 1:
                source_sheet = str(source_sheets[0]).strip()
                source_subjects = dataframe.loc[
                    dataframe["_SOURCE_SHEET"].eq(source_sheet), "SUBJECTCODE"
                ].nunique()
                sheet_hint = source_sheet if source_subjects == 1 else subject_code
            else:
                sheet_hint = subject_code
            sheet_name = safe_sheet_name(sheet_hint, used_sheet_names)
            students.to_excel(writer, sheet_name=sheet_name, index=False)

        workbook = writer.book
        for ws in workbook.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            thin_black = Side(style="thin", color="000000")
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.font = Font(name="Times New Roman", size=8)
                    cell.border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)

            widths = [10, 8.7, 16.5, 8.7, 8.7, 13, 10.5, 8.7, 15.5, 23.7, 8.7, 8.7, 13, 13, 25.5, 20.5]
            for col_idx, width in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(col_idx)].width = width
            for row in ws.iter_rows(min_row=2, min_col=15, max_col=16):
                for cell in row:
                    if cell.value:
                        cell.number_format = "dd-mm-yyyy"

            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            ws.print_title_rows = "1:1"
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_margins.left = 0.25
            ws.page_margins.right = 0.25
            ws.page_margins.top = 0.5
            ws.page_margins.bottom = 0.5

    output.seek(0)
    return output


@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        details = {
            "degree": request.form.get("degree", "").strip(),
            "college_code": request.form.get("college_code", "").strip(),
            "branch_code": request.form.get("branch_code", "").strip(),
            "section": request.form.get("section", "").strip(),
            "regulation": request.form.get("regulation", "").strip(),
            "year": request.form.get("year", "").strip(),
            "semester": request.form.get("semester", "").strip(),
        }

        missing = [label for key, label in [
            ("degree", "Degree"), ("college_code", "CollegeCode"),
            ("branch_code", "BranchCode"), ("section", "Section"),
            ("regulation", "Regulation"), ("year", "Year"),
            ("semester", "Semester")
        ] if not details[key]]
        if missing:
            flash("Please fill in: " + ", ".join(missing), "danger")
            return render_template("index.html", details=details)

        blob_pathname = request.form.get("blob_pathname", "").strip()
        uploaded = request.files.get("excel_file")
        if blob_pathname:
            extension = os.path.splitext(blob_pathname)[1].lower()
            if not blob_pathname.startswith("uploads/") or extension not in {".xlsx", ".xlsm"}:
                flash("The uploaded workbook reference is invalid.", "danger")
                return render_template("index.html", details=details)
        elif not uploaded or not uploaded.filename:
            flash("Please choose an Excel file.", "danger")
            return render_template("index.html", details=details)
        elif not allowed_file(uploaded.filename):
            flash("Upload an .xlsx or .xlsm file. Legacy .xls files are not supported; save them as .xlsx first.", "danger")
            return render_template("index.html", details=details)

        try:
            if blob_pathname:
                workbook_source = BytesIO(download_private_blob(blob_pathname))
            else:
                workbook_source = uploaded
            workbook_data = pd.read_excel(workbook_source, engine="openpyxl", dtype=str, sheet_name=None)
            frames = []
            for sheet_name, sheet_data in workbook_data.items():
                if sheet_data.empty:
                    continue
                sheet_data.columns = [
                    COLUMN_ALIASES.get(str(col).strip().upper(), str(col).strip().upper())
                    for col in sheet_data.columns
                ]
                missing_columns = [col for col in REQUIRED_COLUMNS if col not in sheet_data.columns]
                if missing_columns:
                    flash(
                        f"Sheet '{sheet_name}' is missing required Excel columns: " + ", ".join(missing_columns),
                        "danger"
                    )
                    return render_template("index.html", details=details)
                sheet_data["_SOURCE_SHEET"] = sheet_name
                frames.append(sheet_data)

            if not frames:
                flash("The uploaded Excel file has no student records.", "danger")
                return render_template("index.html", details=details)

            dataframe = pd.concat(frames, ignore_index=True, sort=False)
            for col in REQUIRED_COLUMNS:
                dataframe[col] = dataframe[col].fillna("").astype(str).str.strip()
            for col in OUTPUT_COLUMNS:
                if col not in dataframe.columns:
                    dataframe[col] = ""

            dataframe = dataframe[dataframe["HTNO"].ne("")]
            if dataframe.empty:
                flash("No non-empty HTNO values were found.", "danger")
                return render_template("index.html", details=details)

            if dataframe["FORTNIGHT"].eq("").all():
                flash("No non-empty FORTNIGHT values were found.", "danger")
                return render_template("index.html", details=details)

            # Ensure each group has meaningful subject information.
            if dataframe["SUBJECTCODE"].eq("").all() or dataframe["SUBJECTNAME"].eq("").all():
                flash("SUBJECTCODE and SUBJECTNAME must contain values.", "danger")
                return render_template("index.html", details=details)

            output = make_output(dataframe, details)
            filename = f"Fortnight_Attendance_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            return send_file(
                output,
                as_attachment=True,
                download_name=filename,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as exc:
            app.logger.exception("Failed to process uploaded Excel file")
            flash(f"Could not process the Excel file. Check that it is a valid Excel workbook. Details: {exc}", "danger")
            return render_template("index.html", details=details)
        finally:
            if blob_pathname and blob_uploads_enabled():
                try:
                    delete_blob(blob_pathname)
                except Exception:
                    app.logger.exception("Failed to delete uploaded workbook from Blob storage")

    return render_template("index.html", details={})


@app.route("/sample-template")
def sample_template():
    template_path = os.path.join(app.root_path, "SampleTemplate.xlsx")
    return send_file(
        template_path,
        as_attachment=True,
        download_name="SampleTemplate.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.errorhandler(413)
def file_too_large(_error):
    flash("File is too large. Maximum upload size is 20 MB.", "danger")
    return redirect(url_for("index"))


from consolidated.api.index import app as consolidated_app

app.wsgi_app = DispatcherMiddleware(
    app.wsgi_app,
    {"/consolidated": consolidated_app},
)


if __name__ == "__main__":
    app.run(debug=True)
