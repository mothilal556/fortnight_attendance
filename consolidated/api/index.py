import os
import re
import tempfile
import uuid
import pandas as pd
import openpyxl
from io import BytesIO
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from flask import Flask, render_template, request, send_file, jsonify, url_for
from blob_storage import delete_blob, download_private_blob, upload_private_blob

app = Flask(__name__, template_folder='../templates')
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "local-development-only-secret")
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(os.environ.get("VERCEL")),
)
UPLOAD_FOLDER = tempfile.gettempdir()
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = (
    20 * 1024 * 1024
    if not os.environ.get("VERCEL") or os.environ.get("BLOB_READ_WRITE_TOKEN")
    else 4 * 1024 * 1024
)


def blob_uploads_enabled():
    return bool(os.environ.get("VERCEL") and os.environ.get("BLOB_READ_WRITE_TOKEN"))


@app.context_processor
def inject_blob_upload_setting():
    return {
        "blob_uploads_enabled": blob_uploads_enabled(),
        "upload_limit_mb": app.config["MAX_CONTENT_LENGTH"] // (1024 * 1024),
    }


def validate_sheet_attendance(df, sheet_name, conducted_col, attended_col):
    """Return row-level validation errors when attended exceeds conducted."""
    if not conducted_col or not attended_col:
        return []

    errors = []
    for row_index, row in df[[conducted_col, attended_col]].iterrows():
        conducted_value = pd.to_numeric(row[conducted_col], errors='coerce')
        attended_value = pd.to_numeric(row[attended_col], errors='coerce')

        if pd.isna(conducted_value) or pd.isna(attended_value):
            continue

        if attended_value > conducted_value:
            errors.append(
                f"Sheet '{sheet_name}' row {row_index + 2}: ClassesAttended ({int(attended_value)}) exceeds ClassesConducted ({int(conducted_value)})"
            )

    return errors


def process_excel(file_path, output_path):
    """
    Reads all sheets from the uploaded workbook, sums up student attendance/scores,
    calculates totals and percentages, and returns structured data and file output path.
    """
    excel_file = pd.ExcelFile(file_path)
    sheet_names = excel_file.sheet_names
    excel_file.close()

    attended_dict = {}
    conducted_dict = {}
    validation_errors = []

    # Extract data from each sheet in the workbook
    for sheet in sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet)

        # Clean column names
        df.columns = [str(c).strip().upper() for c in df.columns]

        # Identify student identifier column (HTNO, ROLL NO, PIN, etc.)
        htno_col = None
        for candidate in ['HTNO', 'ROLL NO', 'PIN', 'STUDENT ID', 'REGISTER NO']:
            if candidate in df.columns:
                htno_col = candidate
                break
        if not htno_col:
            htno_col = df.columns[0]  # Fallback to first column

        # Identify attendance / marks conducted column
        conducted_col = None
        for candidate in ['CLASSESCONDUCTED', 'TOTAL CLASSES', 'CONDUCTED', 'MAX MARKS', 'TOTAL MARKS']:
            if candidate in df.columns:
                conducted_col = candidate
                break

        # Identify attendance / marks attended column
        attended_col = None
        for candidate in ['CLASSESATTENDED', 'ATTENDED', 'MARKS', 'SCORE', 'PRESENT']:
            if candidate in df.columns:
                attended_col = candidate
                break

        validation_errors.extend(validate_sheet_attendance(df, sheet, conducted_col, attended_col))

        df[htno_col] = df[htno_col].astype(str).str.strip()

        # Deduce total conducted classes/marks for this subject/sheet
        if conducted_col and not df[conducted_col].empty:
            first_htno = df[htno_col].iloc[0]
            total_conducted = int(df[df[htno_col] == first_htno][conducted_col].sum())
        else:
            total_conducted = 100  # Default fallback benchmark

        conducted_dict[sheet] = total_conducted

        # Aggregate attended values per student ID
        if attended_col:
            grouped = df.groupby(htno_col)[attended_col].sum()
        else:
            val_col = df.columns[1] if len(df.columns) > 1 else df.columns[0]
            grouped = df.groupby(htno_col)[val_col].apply(lambda x: pd.to_numeric(x, errors='coerce').sum())

        for htno, val in grouped.items():
            if htno == 'nan' or not htno:
                continue
            if htno not in attended_dict:
                attended_dict[htno] = {}
            attended_dict[htno][sheet] = int(val) if pd.notnull(val) else 0

    if validation_errors:
        raise ValueError("\n".join(validation_errors))

    # Build consolidated table structure
    rows = []
    
    # Header summary row (Max limits / Conducted per subject)
    top_row = {'HTNO': 'MAX / CONDUCTED'}
    total_max_conducted = sum(conducted_dict.values())
    for sheet in sheet_names:
        top_row[sheet] = conducted_dict[sheet]
    top_row['TOTAL'] = total_max_conducted
    top_row['%'] = '-'
    rows.append(top_row)

    # Student data rows
    for htno, subj_att in attended_dict.items():
        row = {'HTNO': htno}
        total_att = 0
        for sheet in sheet_names:
            att = subj_att.get(sheet, 0)
            row[sheet] = att
            total_att += att
        
        row['TOTAL'] = total_att
        pct = round((total_att / total_max_conducted) * 100, 2) if total_max_conducted > 0 else 0.0
        row['%'] = pct
        rows.append(row)

    consolidated_df = pd.DataFrame(rows)

    # Generate styled Excel file with openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidated Summary"

    headers = list(consolidated_df.columns)
    ws.append(headers)

    for _, r in consolidated_df.iterrows():
        ws.append(list(r.values))

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    max_row_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    
    header_font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
    bold_font = Font(name='Calibri', size=11, bold=True)
    normal_font = Font(name='Calibri', size=11)
    center_align = Alignment(horizontal='center', vertical='center')

    # Apply styling to cells
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_col=len(headers), max_row=len(consolidated_df) + 1), start=2):
        is_max_row = (row_idx == 2)
        for col_idx, cell in enumerate(row, start=1):
            cell.border = thin_border
            cell.alignment = center_align
            
            if is_max_row:
                cell.font = bold_font
                cell.fill = max_row_fill
            else:
                cell.font = normal_font

            if col_idx == len(headers) and not is_max_row:
                cell.number_format = '0.00'

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)

    # Return structured JSON data for frontend live preview
    preview_data = {
        "headers": headers,
        "max_row": rows[0],
        "students": rows[1:],
        "sheet_names": sheet_names,
        "total_students": len(rows) - 1,
        "total_conducted": total_max_conducted
    }

    return output_path, preview_data


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/process', methods=['POST'])
def process():
    blob_pathname = ""
    if 'file' not in request.files:
        if request.is_json:
            blob_pathname = str(request.json.get("blob_pathname", "")).strip()
            if not blob_pathname.startswith("uploads/"):
                return jsonify({'error': 'No valid uploaded workbook reference was provided.'}), 400
            filename = blob_pathname.rsplit("/", 1)[-1]
        else:
            return jsonify({'error': 'No file uploaded'}), 400
    else:
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'No selected file'}), 400
        filename = file.filename

    extension = os.path.splitext(filename)[1].lower()
    if extension in {'.xlsx', '.xls'}:
        upload_path = None
        token = uuid.uuid4().hex
        output_path = os.path.join(app.config['UPLOAD_FOLDER'], f"Consolidated_Report_{token}.xlsx")
        try:
            if blob_pathname:
                with tempfile.NamedTemporaryFile(delete=False, suffix=extension) as upload:
                    upload_path = upload.name
                    upload.write(download_private_blob(blob_pathname))
            else:
                with tempfile.NamedTemporaryFile(delete=False, suffix=extension) as upload:
                    upload_path = upload.name
                    file.save(upload)

            _, preview_data = process_excel(upload_path, output_path)
            if blob_uploads_enabled():
                with open(output_path, "rb") as report:
                    upload_private_blob(
                        f"reports/{token}.xlsx",
                        report.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                os.remove(output_path)
            return jsonify({
                'success': True,
                'preview': preview_data,
                'download_url': url_for('download', token=token)
            })
        except ValueError as e:
            if os.path.exists(output_path):
                os.remove(output_path)
            return jsonify({'error': str(e)}), 400
        except Exception as e:
            app.logger.exception("Failed to process consolidated workbook")
            if os.path.exists(output_path):
                os.remove(output_path)
            return jsonify({'error': f"Processing error: {str(e)}"}), 500
        finally:
            if upload_path and os.path.exists(upload_path):
                os.remove(upload_path)
            if blob_pathname:
                try:
                    delete_blob(blob_pathname)
                except Exception:
                    app.logger.exception("Failed to delete uploaded workbook from Blob storage")
    else:
        if blob_pathname:
            try:
                delete_blob(blob_pathname)
            except Exception:
                app.logger.exception("Failed to delete unsupported uploaded workbook from Blob storage")
        return jsonify({'error': 'Invalid file type. Please upload an Excel workbook (.xlsx or .xls)'}), 400


@app.route('/download')
def download():
    token = request.args.get('token', '')
    if not re.fullmatch(r'[0-9a-f]{32}', token):
        return jsonify({'error': 'Invalid or missing report token.'}), 400

    output_path = os.path.join(
        app.config['UPLOAD_FOLDER'],
        f"Consolidated_Report_{token}.xlsx"
    )
    if os.path.exists(output_path):
        response = send_file(output_path, as_attachment=True, download_name="Consolidated_Report.xlsx")
        response.call_on_close(lambda: os.path.exists(output_path) and os.remove(output_path))
        return response
    if blob_uploads_enabled():
        blob_pathname = f"reports/{token}.xlsx"
        try:
            report_data = download_private_blob(blob_pathname)
        except Exception:
            app.logger.exception("Failed to retrieve consolidated report from Blob storage")
            return jsonify({'error': 'The report could not be retrieved. Please process the workbook again.'}), 500

        try:
            delete_blob(blob_pathname)
        except Exception:
            app.logger.exception("Failed to delete downloaded report from Blob storage")

        return send_file(
            BytesIO(report_data),
            as_attachment=True,
            download_name="Consolidated_Report.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    return jsonify({'error': 'The report is no longer available. Please process the workbook again.'}), 404


if __name__ == '__main__':
    app.run(debug=True, port=5000)